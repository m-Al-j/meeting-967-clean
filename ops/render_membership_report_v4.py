from __future__ import annotations
import copy
import json
import re
import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEAM_DISPLAY = [
    {"key":"governance","label":"الإدارة والحوكمة","matches":["الاداره والحوكمة","الإدارة والحوكمة"]},
    {"key":"hr","label":"الموارد البشرية","matches":["الموارد البشرية"]},
    {"key":"projects","label":"المشاريع والخدمات المجتمعية","matches":["الفريق التنفيذي","المشاريع والخدمات المجتمعية"]},
    {"key":"tech","label":"التقنية والبحث والبيانات","matches":["التقنية والبحث","التقنية والبحث والبيانات"]},
    {"key":"media","label":"الإعلام والعلاقات","matches":["الفريق الإعلامي","الإعلام والعلاقات"]},
    {"key":"other","label":"أخرى / عضوية عامة","matches":[]},
]

ARABIC_REPLACEMENTS = str.maketrans({'أ':'ا','إ':'ا','آ':'ا','ة':'ه'})
ARABIC_RE = re.compile(r'[\u0600-\u06ff]')


def clean(v, max_len=500):
    return str(v if v is not None else '').strip()[:max_len]


def parse_array(v):
    if isinstance(v, list):
        return [str(x) for x in v]
    try:
        x = json.loads(str(v or '[]'))
        return [str(a) for a in x] if isinstance(x, list) else []
    except Exception:
        return []


def norm(v):
    return clean(v).lower().translate(ARABIC_REPLACEMENTS).replace(' ','').replace('\u200f','').replace('\u200e','')


def team_of(names):
    n = {norm(x) for x in names}
    for t in TEAM_DISPLAY[:-1]:
        if any(norm(x) in n for x in t['matches']):
            return t
    return TEAM_DISPLAY[-1]


def contains_arabic(text):
    return bool(ARABIC_RE.search(str(text)))


def set_direction(paragraph, rtl: bool):
    """Use paragraph bidi only. Do not set w:rtl on runs; it can break Arabic shaping in some viewers."""
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn('w:bidi'))
    if rtl:
        if bidi is None:
            bidi = OxmlElement('w:bidi')
            ppr.append(bidi)
        bidi.set(qn('w:val'), '1')
    elif bidi is not None:
        ppr.remove(bidi)

    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rtl_el = rpr.find(qn('w:rtl'))
        if rtl_el is not None:
            rpr.remove(rtl_el)

        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.insert(0, rfonts)
        rfonts.set(qn('w:ascii'), 'Noto Sans Arabic')
        rfonts.set(qn('w:hAnsi'), 'Noto Sans Arabic')
        rfonts.set(qn('w:cs'), 'Noto Sans Arabic')

        lang = rpr.find(qn('w:lang'))
        if lang is None:
            lang = OxmlElement('w:lang')
            rpr.append(lang)
        lang.set(qn('w:val'), 'ar-SA' if rtl else 'en-US')
        lang.set(qn('w:bidi'), 'ar-SA' if rtl else 'en-US')


def clear_paragraph(paragraph):
    for run in list(paragraph.runs):
        run.text = ''
    if not paragraph.runs:
        paragraph.add_run('')


def set_cell_text(cell, text, align=None, rtl=None):
    text = clean(text, 500)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    if rtl is None:
        rtl = contains_arabic(text)
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)
    set_direction(p, rtl)


def set_meta_value(cell, text, rtl=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text = clean(text)
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)
    set_direction(p, contains_arabic(text) if rtl is None else rtl)


def clone_row(table, template_row_index=1, template_tr=None):
    tr = copy.deepcopy(template_tr if template_tr is not None else table.rows[template_row_index]._tr)
    table._tbl.append(tr)
    return table.rows[-1]


def remove_rows_after(table, keep):
    while len(table.rows) > keep:
        table._tbl.remove(table.rows[-1]._tr)


def repeat_header(table):
    if not table.rows:
        return
    first = table.rows[0]._tr.get_or_add_trPr()
    if first.find(qn('w:tblHeader')) is None:
        first.append(OxmlElement('w:tblHeader'))
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn('w:cantSplit')) is None:
            trpr.append(OxmlElement('w:cantSplit'))


def local_text(value):
    return clean(value) if value else '—'


def replace_heading(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == 'تقرير إدارة العضوية':
            clear_paragraph(p)
            p.runs[0].text = text
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_direction(p, True)
            return



def main(payload_path, output_path):
    payload = json.loads(
        Path(payload_path).read_text(encoding='utf-8')
    )

    doc = Document(payload['template'])
    rows = payload.get('rows') or []
    stats = payload.get('stats') or {}

    withdrawn = [
        r for r in rows
        if r.get('choice') == 'withdraw'
    ]

    frozen = [
        r for r in rows
        if r.get('choice') == 'freeze'
    ]

    continuing = [
        r for r in rows
        if r.get('choice') == 'continue'
    ]

    no_response = [
        r for r in rows
        if r.get('choice') in (None, '')
    ]

    # ---------------------------------
    # توزيع الأعضاء حسب الفرق
    # ---------------------------------
    distribution = []

    for t in TEAM_DISPLAY:
        distribution.append({
            **t,
            'total': 0,
            'continue': 0,
            'freeze': 0,
            'withdraw': 0
        })

    for r in rows:
        t = team_of(
            parse_array(r.get('team_role_names'))
        )

        dst = next(
            x for x in distribution
            if x['key'] == t['key']
        )

        dst['total'] += 1

        if r.get('choice') == 'continue':
            dst['continue'] += 1
        elif r.get('choice') == 'freeze':
            dst['freeze'] += 1
        elif r.get('choice') == 'withdraw':
            dst['withdraw'] += 1

    # ---------------------------------
    # البيانات التعريفية
    # ---------------------------------
    meta = doc.tables[1].rows[0].cells

    set_meta_value(
        meta[0],
        f"[ {payload['report_id']} ]",
        rtl=False
    )

    set_meta_value(
        meta[1],
        f"[ {payload['period_start']} — {payload['period_end']} ]",
        rtl=False
    )

    set_meta_value(
        meta[2],
        f"[ {payload['issue_date']} ]",
        rtl=False
    )

    set_meta_value(
        meta[3],
        'فريق الموارد البشرية',
        rtl=True
    )

    # ---------------------------------
    # بطاقات الإحصائيات
    # إزالة أي رقم ثانوي قديم
    # ---------------------------------
    cards = doc.tables[2].rows[0].cells

    card_values = [
        stats.get('total', 0),
        stats.get('continue', 0),
        stats.get('freeze', 0),
        stats.get('withdraw', 0)
    ]

    for cell, value in zip(cards, card_values):
        paragraphs = list(cell.paragraphs)

        if len(paragraphs) >= 2:
            value_paragraph = paragraphs[1]
        else:
            value_paragraph = cell.add_paragraph()

        clear_paragraph(value_paragraph)
        value_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_paragraph.add_run(str(value))
        set_direction(value_paragraph, False)

        for extra in paragraphs[2:]:
            parent = extra._element.getparent()
            if parent is not None:
                parent.remove(extra._element)

    # ---------------------------------
    # التوزيع العام
    # ---------------------------------
    dt = doc.tables[3]

    for idx, team in enumerate(distribution, start=1):
        row = dt.rows[idx]

        values = [
            team['label'],
            str(team['total']),
            str(team['continue']),
            str(team['freeze']),
            str(team['withdraw'])
        ]

        for j, value in enumerate(values):
            set_cell_text(
                row.cells[j],
                value,
                WD_ALIGN_PARAGRAPH.CENTER
                if j
                else WD_ALIGN_PARAGRAPH.RIGHT,
                rtl=(j == 0)
            )

    total_row = dt.rows[7]

    values = [
        'المجموع',
        str(stats.get('total', 0)),
        str(stats.get('continue', 0)),
        str(stats.get('freeze', 0)),
        str(stats.get('withdraw', 0))
    ]

    for j, value in enumerate(values):
        set_cell_text(
            total_row.cells[j],
            value,
            WD_ALIGN_PARAGRAPH.CENTER
            if j
            else WD_ALIGN_PARAGRAPH.RIGHT,
            rtl=(j == 0)
        )

    repeat_header(dt)

    # ---------------------------------
    # دالة ملء جدول جاهز
    # ---------------------------------
    def fill_table(table, records, empty_values):
        template_row = (
            copy.deepcopy(table.rows[1]._tr)
            if len(table.rows) > 1
            else None
        )

        remove_rows_after(table, 1)

        if not records:
            row = clone_row(
                table,
                1,
                template_row
            )

            for j, value in enumerate(empty_values):
                set_cell_text(
                    row.cells[j],
                    value,
                    WD_ALIGN_PARAGRAPH.CENTER,
                    rtl=contains_arabic(value)
                )

            repeat_header(table)
            return

        for values in records:
            row = clone_row(
                table,
                1,
                template_row
            )

            for j, value in enumerate(values):
                value = local_text(value)
                rtl = contains_arabic(value)

                if j == 0:
                    align = WD_ALIGN_PARAGRAPH.CENTER
                elif rtl:
                    align = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    align = WD_ALIGN_PARAGRAPH.LEFT

                set_cell_text(
                    row.cells[j],
                    value,
                    align,
                    rtl
                )

        repeat_header(table)

    # ---------------------------------
    # المنسحبون
    #
    # ترتيب الخلايا مطابق للقالب:
    # # | الاسم | رقم العضوية | الفريق | تاريخ الرد | السبب
    #
    # البيانات مأخوذة من snapshot الحملة
    # قبل خروج العضو من السيرفر.
    # ---------------------------------
    withdrawn_records = []

    for idx, r in enumerate(withdrawn, start=1):
        display_name = (
            r.get('display_name')
            or r.get('username')
            or r.get('user_id')
            or '—'
        )

        username = r.get('username') or ''

        if username:
            name_value = f"{display_name}\n@{username}"
        else:
            name_value = display_name

        team_value = team_of(
            parse_array(r.get('team_role_names'))
        )['label']

        withdrawn_records.append([
            str(idx).zfill(2),
            name_value,
            r.get('user_id') or '—',
            team_value,
            r.get('responded_at') or '—',
            r.get('reason')
            or r.get('withdraw_reason')
            or '—'
        ])

    fill_table(
        doc.tables[4],
        withdrawn_records,
        [
            '—',
            'لا توجد حالات انسحاب مسجلة في هذا التقرير',
            '—',
            '—',
            '—',
            '—'
        ]
    )

    # ---------------------------------
    # المجمدون
    #
    # # | الاسم | الفريق | بداية التجميد | تاريخ العودة
    # ---------------------------------
    frozen_records = []

    for idx, r in enumerate(frozen, start=1):
        display_name = (
            r.get('display_name')
            or r.get('username')
            or r.get('user_id')
            or '—'
        )

        frozen_records.append([
            str(idx).zfill(2),
            display_name,
            team_of(
                parse_array(r.get('team_role_names'))
            )['label'],
            r.get('freeze_start_at') or '—',
            r.get('return_at') or '—'
        ])

    fill_table(
        doc.tables[5],
        frozen_records,
        [
            '—',
            'لا توجد حالات تجميد مسجلة في هذا التقرير',
            '—',
            '—',
            '—'
        ]
    )

    # ---------------------------------
    # جدول المستمرين
    # # | الاسم | الفريق | التاريخ | الحالة
    # ---------------------------------
    continuing_table = doc.tables[6]

    continuing_template = (
        copy.deepcopy(continuing_table.rows[1]._tr)
        if len(continuing_table.rows) > 1
        else None
    )

    # نأخذ نسخة مستقلة لإنشاء جدول "لم يردوا"
    pending_table_xml = copy.deepcopy(
        continuing_table._tbl
    )

    # heading الخاص بالقسم الحالي
    section_heading = None

    for paragraph in doc.paragraphs:
        if 'المستمرون والاستجابات المتبقية' in paragraph.text:
            section_heading = paragraph
            clear_paragraph(paragraph)
            paragraph.add_run('المستمرون')
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_direction(paragraph, True)
            break

    remove_rows_after(
        continuing_table,
        1
    )

    for idx, r in enumerate(continuing, start=1):
        row = clone_row(
            continuing_table,
            1,
            continuing_template
        )

        display_name = (
            r.get('display_name')
            or r.get('username')
            or r.get('user_id')
            or '—'
        )

        values = [
            str(idx).zfill(2),
            display_name,
            team_of(
                parse_array(r.get('team_role_names'))
            )['label'],
            r.get('responded_at') or '—',
            'مستمر'
        ]

        for j, value in enumerate(values):
            value = local_text(value)
            rtl = contains_arabic(value)

            if j in (0, 3, 4):
                align = WD_ALIGN_PARAGRAPH.CENTER
            elif rtl:
                align = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                align = WD_ALIGN_PARAGRAPH.LEFT

            set_cell_text(
                row.cells[j],
                value,
                align,
                rtl
            )

    repeat_header(continuing_table)

    # ---------------------------------
    # جدول منفصل للأشخاص الذين لم يردوا
    # ---------------------------------
    if section_heading is not None:
        pending_heading = copy.deepcopy(
            section_heading._p
        )

        clear_paragraph(
            type(
                'ParagraphProxy',
                (),
                {'runs': []}
            )()
        ) if False else None

        # إنشاء نص العنوان مباشرة على الـ XML
        for child in list(pending_heading):
            pending_heading.remove(child)

        ppr = OxmlElement('w:pPr')
        pending_heading.append(ppr)

        bidi = OxmlElement('w:bidi')
        bidi.set(qn('w:val'), '1')
        ppr.append(bidi)

        run = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')

        rfonts = OxmlElement('w:rFonts')
        rfonts.set(qn('w:ascii'), 'Noto Sans Arabic')
        rfonts.set(qn('w:hAnsi'), 'Noto Sans Arabic')
        rfonts.set(qn('w:cs'), 'Noto Sans Arabic')
        rpr.append(rfonts)

        run.append(rpr)

        text = OxmlElement('w:t')
        text.text = f'لم يردوا ({len(no_response)})'
        run.append(text)

        pending_heading.append(run)

        # نسخ جدول المستمرين
        pending_table = pending_table_xml

        # وضع النسخة بعد جدول المستمرين
        continuing_table._tbl.addnext(
            pending_heading
        )

        pending_heading.addnext(
            pending_table
        )

        # بعد الإدراج نستخدم آخر نسخة من الجدول
        pending_table_obj = doc.tables[-1]

        pending_template = (
            copy.deepcopy(
                pending_table_obj.rows[1]._tr
            )
            if len(pending_table_obj.rows) > 1
            else None
        )

        remove_rows_after(
            pending_table_obj,
            1
        )

        for idx, r in enumerate(no_response, start=1):
            row = clone_row(
                pending_table_obj,
                1,
                pending_template
            )

            display_name = (
                r.get('display_name')
                or r.get('username')
                or r.get('user_id')
                or '—'
            )

            values = [
                str(idx).zfill(2),
                display_name,
                team_of(
                    parse_array(r.get('team_role_names'))
                )['label'],
                '—',
                'لم يرد بعد'
            ]

            for j, value in enumerate(values):
                value = local_text(value)
                rtl = contains_arabic(value)

                if j in (0, 3, 4):
                    align = WD_ALIGN_PARAGRAPH.CENTER
                elif rtl:
                    align = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    align = WD_ALIGN_PARAGRAPH.LEFT

                set_cell_text(
                    row.cells[j],
                    value,
                    align,
                    rtl
                )

        repeat_header(pending_table_obj)

    # ---------------------------------
    # العنوان الرئيسي
    # ---------------------------------
    report_title = clean(
        payload.get('report_title')
        or 'تقرير إدارة العضوية — إعلان التصفية',
        180
    )

    replace_heading(
        doc,
        report_title
    )

    # ---------------------------------
    # العنوان الفرعي الجديد
    # ---------------------------------
    for paragraph in doc.paragraphs:
        if (
            'نموذج مؤسسي شامل لمراجعة وتنظيم حالات العضوية'
            in paragraph.text
        ):
            clear_paragraph(paragraph)
            paragraph.add_run(
                'تقرير شامل لمراجعة وتنظيم حالات العضوية — مبادرة 967'
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_direction(paragraph, True)

    # ---------------------------------
    # حذف ملاحظات التقرير بالكامل
    # ---------------------------------
    for paragraph in list(doc.paragraphs):
        if 'ملاحظات التقرير:' in paragraph.text:
            parent = paragraph._element.getparent()
            if parent is not None:
                parent.remove(paragraph._element)

    # ---------------------------------
    # تنظيف نهائي
    # ---------------------------------
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                )

                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        set_direction(
                            paragraph,
                            contains_arabic(paragraph.text)
                        )

    doc.core_properties.title = report_title
    doc.core_properties.subject = (
        'تقرير إدارة العضوية — إعلان التصفية'
    )
    doc.core_properties.author = (
        'فريق الموارد البشرية — مبادرة 967'
    )

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)

    print(out)

if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit('usage: render_membership_report_v4.py PAYLOAD.json OUTPUT.docx')
    main(sys.argv[1], sys.argv[2])
